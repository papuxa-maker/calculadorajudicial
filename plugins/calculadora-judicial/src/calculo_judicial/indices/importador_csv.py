"""Carga manual de índices a partir de CSV/XLSX (tabelas de tribunal e séries sem API).

Formatos aceitos (SPEC, seção 5):

- ``mes_referencia,taxa_pct`` — taxa mensal em %; ou
- ``mes_referencia,fator`` — fator/coeficiente acumulado; a taxa é derivada.

Datas aceitas em ``mes_referencia``: ``aaaa-mm-dd``, ``aaaa-mm``, ``mm/aaaa``,
``dd/mm/aaaa`` ou ``jan/2024`` (sempre normalizadas para o dia 01). CSV em UTF-8
(ou cp1252, o "CSV separado por ponto e vírgula" do Excel em português),
separador vírgula com decimal ponto, ou separador ponto e vírgula com decimal
vírgula. Coluna opcional ``projecao`` (sim/não, 1/0, verdadeiro/falso).
Células XLSX em formato percentual (``0,43%`` guardado como 0,0043) são
convertidas para 0,43.

Derivação da taxa a partir do fator:

- ``fator_fim_do_mes``: o fator do mês M é o nível depois de aplicada a taxa de M;
  ``taxa(M) = fator(M) / fator(M-1) - 1``; o primeiro mês do arquivo não tem taxa.
- ``fator_inicio_do_mes`` (TJSP: "o fator do mês M embute o índice de M−1"):
  ``taxa(M) = fator(M+1) / fator(M) - 1``; o último mês do arquivo não tem taxa.
- ``coeficiente_multiplicador`` (CJF/SICOM: "valor da época × coeficiente"):
  ``taxa(M) = coef(M) / coef(M+1) - 1``; o último mês do arquivo não tem taxa.

Validação (a mesma regra do coletor, SPEC seção 5): taxa dentro da faixa do
índice (−2% a +5%; IGP-M −3% a +8%), meses em sequência sem lacuna (inclusive
em relação ao que já está no banco), sem duplicata. A importação é atômica:
se qualquer mês falhar, nada é gravado e o motivo vai para ``log/alertas.log``.
Tabelas oficiais com meses de inflação alta (ex.: IPC-r de 1994) exigem
``aceitar_fora_da_faixa=True``, e cada mês aceito assim é registrado no log.

Importar sobre um código que hoje é tabela composta exige ``substituir=True``, exceto as
tabelas montadas pela calculadora (TJGO e as dos tribunais, marca "montada pela calculadora"):
essas são trocadas automaticamente pela tabela importada (com registro no log) quando o arquivo
cobre desde o início da cadeia; se o arquivo começa depois, a troca apagaria meses e exige
``substituir=True``. Com 12+ meses, fatores que diminuem (fator) ou coeficientes que aumentam
(coeficiente multiplicador) indicam convenção trocada e o arquivo é recusado.
"""

from __future__ import annotations

import csv
import json
import sqlite3
from dataclasses import dataclass, field
from datetime import date
from decimal import Decimal, localcontext
from pathlib import Path

from calculo_judicial.datas import parse_data, primeiro_dia, somar_meses
from calculo_judicial.decimais import CEM, CONTEXTO, UM, para_decimal, q10
from calculo_judicial.indices.fatores import recalcular_fatores
from calculo_judicial.indices.modelo import InfoIndice, gravar_taxa, registrar_indice

CONVENCOES_FATOR = ("fator_fim_do_mes", "fator_inicio_do_mes", "coeficiente_multiplicador")
MINIMO_PARA_TENDENCIA = 12  # razões mensais para julgar se os fatores crescem (fator) ou diminuem (coeficiente)


def _pct_br(valor: Decimal) -> str:
    return f"{valor:.4f}".replace(".", ",") + "%"
VERDADEIROS = {"1", "1.0", "sim", "s", "true", "verdadeiro", "v", "x", "yes", "y"}
FALSOS = {"", "0", "0.0", "nao", "não", "n", "false", "falso", "f", "no"}
MARCA_PRESUMIDA = "montada pela calculadora"  # cadeia montada pela calculadora (não é tabela oficial importada)
# as outras: nomes usados nas primeiras versões (bancos já instalados)
MARCAS_PRESUMIDA = (MARCA_PRESUMIDA, "confira a tabela do processo", "pendente de confirmação")


class ErroImportacao(ValueError):
    """Arquivo recusado: nada foi gravado."""


@dataclass
class ResultadoImportacao:
    indice: str
    arquivo: str
    gravados: int = 0
    substituidos: int = 0
    ignorados: list[str] = field(default_factory=list)
    alertas: list[str] = field(default_factory=list)
    primeiro_mes: date | None = None
    ultimo_mes: date | None = None
    modo: str = ""  # "taxa" (coluna taxa_pct) ou "fator"
    meses_do_arquivo: list[date] = field(default_factory=list)  # meses com taxa derivada do arquivo

    def resumo(self) -> str:
        faixa = ""
        if self.primeiro_mes and self.ultimo_mes:
            faixa = f" de {self.primeiro_mes:%m/%Y} a {self.ultimo_mes:%m/%Y}"
        texto = f"{self.indice}: {self.gravados} meses gravados{faixa} a partir de {self.arquivo}"
        if self.substituidos:
            texto += f" ({self.substituidos} substituídos)"
        if self.ignorados:
            n = len(self.ignorados)
            texto += f"; {n} {'linha não gravada' if n == 1 else 'linhas não gravadas'}: " + "; ".join(self.ignorados[:5])
        if self.alertas:
            texto += "; avisos: " + "; ".join(self.alertas[:4]) + (f" (e mais {len(self.alertas) - 4})" if len(self.alertas) > 4 else "")
        return texto


@dataclass
class _Linha:
    numero: int
    mes: date
    valor: Decimal
    projecao: bool


def _normalizar_coluna(nome: object) -> str:
    n = str(nome).strip().lower()
    for a, b in (("ã", "a"), ("ç", "c"), ("ê", "e"), ("é", "e"), ("ó", "o"), ("í", "i"), ("á", "a"), ("ê", "e")):
        n = n.replace(a, b)
    return n.replace(" ", "_")


def _ler_projecao(valor: object, numero: int) -> bool:
    if isinstance(valor, bool):
        return valor
    texto = str(valor if valor is not None else "").strip().lower()
    if texto in VERDADEIROS:
        return True
    if texto in FALSOS:
        return False
    raise ErroImportacao(f"linha {numero}: valor de projecao não reconhecido: {valor!r}")


def ler_linhas(caminho: str | Path) -> list[dict[str, object]]:
    """Lê CSV ou XLSX e devolve dicionários com colunas normalizadas.

    Valores de XLSX numéricos chegam como texto (``repr`` do float, que reproduz o
    número digitado); células em formato percentual são multiplicadas por 100.
    Linhas de CSV com campos a mais ou a menos que o cabeçalho levantam
    ``ErroImportacao`` (evita gravar "0,43" lido como "0").
    """
    caminho = Path(caminho)
    if caminho.suffix.lower() in (".xlsx", ".xlsm"):
        from openpyxl import load_workbook

        wb = load_workbook(caminho, data_only=True)
        try:
            ws = wb.worksheets[0]
            linhas = list(ws.iter_rows())
            if not linhas:
                return []
            cab = [_normalizar_coluna(c.value) for c in linhas[0]]
            saida: list[dict[str, object]] = []
            for linha in linhas[1:]:
                if all(c.value is None or str(c.value).strip() == "" for c in linha):
                    continue
                registro: dict[str, object] = {}
                for nome, celula in zip(cab, linha):
                    valor = celula.value
                    if isinstance(valor, bool):
                        registro[nome] = valor
                        continue
                    if isinstance(valor, (int, float)) and not isinstance(valor, bool):
                        texto = repr(valor) if isinstance(valor, float) else str(valor)
                        numero = para_decimal(texto)
                        if "%" in str(celula.number_format):
                            numero = numero * CEM
                        registro[nome] = format(numero, "f")
                        continue
                    if hasattr(valor, "isoformat"):
                        valor = valor.isoformat()[:10]
                    registro[nome] = "" if valor is None else str(valor)
                saida.append(registro)
            return saida
        finally:
            wb.close()
    bruto = caminho.read_bytes()
    try:
        texto = bruto.decode("utf-8-sig")
    except UnicodeDecodeError:
        texto = bruto.decode("cp1252")
    linhas_texto = texto.splitlines()
    if not linhas_texto:
        return []
    cabecalho = linhas_texto[0]
    separador = ";" if cabecalho.count(";") > cabecalho.count(",") else ","
    leitor = csv.reader(linhas_texto, delimiter=separador)
    cab = [_normalizar_coluna(c) for c in next(leitor)]
    saida = []
    for numero, campos in enumerate(leitor, start=2):
        if not campos or all(c.strip() == "" for c in campos):
            continue
        if len(campos) != len(cab):
            raise ErroImportacao(
                f"linha {numero}: {len(campos)} campos para {len(cab)} colunas ({separador!r} como separador). "
                "Se o decimal é vírgula, use ponto e vírgula como separador."
            )
        saida.append({nome: valor.strip() for nome, valor in zip(cab, campos)})
    return saida


def _faixa_do_indice(indice: str) -> tuple[Decimal, Decimal]:
    from calculo_judicial.indices import coletor_bcb

    for definicao in coletor_bcb.SERIES_BCB:
        if definicao.codigo == indice:
            return definicao.faixa
    return coletor_bcb.FAIXA_PADRAO


def _alertar(indice: str, mensagem: str, caminho_alertas: Path | None, resultado: ResultadoImportacao) -> None:
    from calculo_judicial.indices.coletor_bcb import registrar_alerta

    registrar_alerta(indice, mensagem, caminho_alertas)
    resultado.alertas.append(mensagem)


def _recusar(indice: str, mensagem: str, caminho_alertas: Path | None) -> None:
    from calculo_judicial.indices.coletor_bcb import registrar_alerta

    registrar_alerta(indice, f"importação recusada: {mensagem}", caminho_alertas)
    raise ErroImportacao(mensagem)


def importar_arquivo(
    conn: sqlite3.Connection,
    caminho: str | Path,
    indice: str,
    nome: str | None = None,
    fonte: str | None = None,
    url_fonte: str | None = None,
    tipo: str = "correcao",
    convencao_fator: str = "fator_fim_do_mes",
    substituir: bool = False,
    aceitar_fora_da_faixa: bool = False,
    desde: date | None = None,
    caminho_alertas: Path | None = None,
) -> ResultadoImportacao:
    """Importa ``caminho`` para o índice ``indice`` e recalcula o fator acumulado.

    Tudo ou nada: o arquivo é lido e validado inteiro antes de qualquer gravação.
    ``desde``: ignora meses anteriores (ex.: tabelas que começam em 1964).
    """
    if convencao_fator not in CONVENCOES_FATOR:
        raise ValueError(f"convenção de fator inválida: {convencao_fator}")
    caminho = Path(caminho)
    resultado = ResultadoImportacao(indice, caminho.name)

    # 1) Leitura e conversão (sem tocar no banco)
    try:
        linhas = ler_linhas(caminho)
    except ErroImportacao as erro:
        _recusar(indice, f"{caminho.name}: {erro}", caminho_alertas)
    if not linhas:
        _recusar(indice, f"{caminho.name}: arquivo vazio", caminho_alertas)
    colunas = set(linhas[0].keys())
    if "mes_referencia" not in colunas:
        _recusar(indice, f"{caminho.name}: coluna mes_referencia ausente; colunas: {sorted(colunas)}", caminho_alertas)
    if "taxa_pct" in colunas:
        modo = "taxa"
    elif "fator" in colunas:
        modo = "fator"
    else:
        _recusar(indice, f"{caminho.name}: precisa da coluna taxa_pct ou fator; colunas: {sorted(colunas)}", caminho_alertas)

    registros: list[_Linha] = []
    erros: list[str] = []
    for numero, linha in enumerate(linhas, start=2):
        try:
            mes = primeiro_dia(parse_data(str(linha["mes_referencia"])))
            valor = para_decimal(str(linha["fator" if modo == "fator" else "taxa_pct"]))
            projecao = _ler_projecao(linha.get("projecao", ""), numero)
        except (ValueError, TypeError) as erro:
            erros.append(f"linha {numero}: {erro}")
            continue
        if desde and mes < desde:
            continue
        registros.append(_Linha(numero, mes, valor, projecao))
    if erros:
        _recusar(indice, f"{caminho.name}: {len(erros)} linhas inválidas: " + "; ".join(erros[:5]), caminho_alertas)
    if not registros:
        _recusar(indice, f"{caminho.name}: nenhuma linha válida", caminho_alertas)
    registros.sort(key=lambda r: r.mes)
    for anterior, atual in zip(registros, registros[1:]):
        if atual.mes == anterior.mes:
            _recusar(indice, f"{caminho.name}: mês duplicado {atual.mes:%m/%Y} (linhas {anterior.numero} e {atual.numero})", caminho_alertas)
        if atual.mes != somar_meses(anterior.mes, 1):
            _recusar(indice, f"{caminho.name}: lacuna entre {anterior.mes:%m/%Y} e {atual.mes:%m/%Y}", caminho_alertas)

    # 2) Taxas mensais
    taxas: list[tuple[date, Decimal, bool]] = []
    if modo == "taxa":
        taxas = [(r.mes, r.valor, r.projecao) for r in registros]
    else:
        if any(r.valor <= 0 for r in registros):
            _recusar(indice, f"{caminho.name}: fator zero ou negativo", caminho_alertas)
        with localcontext(CONTEXTO):
            razoes = [atual.valor / anterior.valor for anterior, atual in zip(registros, registros[1:])]
            for (anterior, atual), razao in zip(zip(registros, registros[1:]), razoes):
                projecao = anterior.projecao or atual.projecao
                if convencao_fator == "fator_fim_do_mes":
                    taxas.append((atual.mes, q10((razao - UM) * CEM), projecao))
                elif convencao_fator == "fator_inicio_do_mes":
                    taxas.append((anterior.mes, q10((razao - UM) * CEM), projecao))
                else:
                    taxas.append((anterior.mes, q10((UM / razao - UM) * CEM), projecao))
        # Tendência dos fatores (só com 12+ meses: um período curto de deflação não decide nada)
        if len(razoes) >= MINIMO_PARA_TENDENCIA:
            decrescentes = sum(1 for r in razoes if r < UM)
            crescentes = sum(1 for r in razoes if r > UM)
            if convencao_fator != "coeficiente_multiplicador" and decrescentes * 2 > len(razoes):
                _recusar(
                    indice,
                    f"{caminho.name}: a maioria dos fatores diminui mês a mês; parece tabela de coeficiente "
                    "multiplicador (use --convencao coeficiente_multiplicador)",
                    caminho_alertas,
                )
            if convencao_fator == "coeficiente_multiplicador" and crescentes * 2 > len(razoes):
                _recusar(
                    indice,
                    f"{caminho.name}: a maioria dos coeficientes aumenta mês a mês; parece fator acumulado "
                    "(use --convencao fator_fim_do_mes ou fator_inicio_do_mes)",
                    caminho_alertas,
                )
        if convencao_fator == "fator_fim_do_mes":
            resultado.ignorados.append(f"{registros[0].mes:%m/%Y}: primeiro fator do arquivo (sem taxa derivável)")
        else:
            resultado.ignorados.append(f"{registros[-1].mes:%m/%Y}: último fator do arquivo (sem taxa derivável)")
    if not taxas:
        _recusar(indice, f"{caminho.name}: nenhuma taxa derivável", caminho_alertas)
    resultado.modo = modo
    resultado.meses_do_arquivo = [m for m, _, _ in taxas]

    # 3) Validação de faixa
    minimo, maximo = _faixa_do_indice(indice)
    fora = [(m, t) for m, t, _ in taxas if not (minimo <= t <= maximo)]
    if fora and not aceitar_fora_da_faixa:
        lista = "; ".join(f"{m:%m/%Y}={_pct_br(t)}" for m, t in fora[:8])
        _recusar(
            indice,
            f"{caminho.name}: {len(fora)} meses fora da faixa usual ({minimo}% a {maximo}% ao mês): {lista}. "
            "Inflação alta de 1994 é normal; se o arquivo é mesmo esse, use --aceitar-fora-da-faixa",
            caminho_alertas,
        )

    # 4) Conferência com o que já está no banco
    composta = conn.execute("SELECT 1 FROM tabelas_compostas WHERE nome = ?", (indice,)).fetchone()
    info_existente = conn.execute("SELECT nome FROM indices WHERE codigo = ?", (indice,)).fetchone()
    composta_presumida = bool(composta and info_existente and any(m in (info_existente[0] or "") for m in MARCAS_PRESUMIDA))
    if composta and not (substituir or composta_presumida):
        _recusar(indice, f"{indice} hoje é tabela composta; para trocá-la pela tabela importada use --substituir", caminho_alertas)
    if composta and composta_presumida and not substituir:
        # a tabela montada pela calculadora só é trocada sem --substituir se o arquivo cobre o mesmo começo
        inicio_cadeia = min(date.fromisoformat(p["de"])
                            for p in json.loads(conn.execute("SELECT periodos_json FROM tabelas_compostas WHERE nome = ?",
                                                             (indice,)).fetchone()[0]))
        primeiro_arquivo = min(m for m, _, _ in taxas)
        if primeiro_arquivo > primeiro_dia(inicio_cadeia):
            _recusar(
                indice,
                f"{indice} hoje é a tabela montada pela calculadora, que cobre desde {inicio_cadeia:%m/%Y}; o arquivo começa "
                f"em {primeiro_arquivo:%m/%Y}, e os meses de {inicio_cadeia:%m/%Y} a {somar_meses(primeiro_arquivo, -1):%m/%Y} "
                "deixariam de existir. Use outro código (ex.: "
                f"{indice}_OFICIAL) ou, para trocar mesmo assim, --substituir",
                caminho_alertas,
            )
    existentes = {
        parse_data(m): (Decimal(t), bool(p))
        for m, t, p in conn.execute(
            "SELECT mes_referencia, taxa_pct, projecao FROM taxas_mensais WHERE indice_codigo = ?", (indice,)
        )
    }
    divergentes = [
        (m, existentes[m][0], t)
        for m, t, proj in taxas
        if m in existentes and not existentes[m][1] and existentes[m][0] != t and not proj
    ]
    if divergentes and not substituir:
        lista = "; ".join(f"{m:%m/%Y}: banco {a}, arquivo {b}" for m, a, b in divergentes[:5])
        _recusar(indice, f"{caminho.name}: {len(divergentes)} meses divergem do banco ({lista}); use --substituir", caminho_alertas)
    meses_finais = set(existentes) | {m for m, _, _ in taxas}
    ordenados = sorted(meses_finais)
    for anterior, atual in zip(ordenados, ordenados[1:]):
        if atual != somar_meses(anterior, 1):
            _recusar(indice, f"{caminho.name}: a série ficaria com lacuna entre {anterior:%m/%Y} e {atual:%m/%Y}", caminho_alertas)

    # 5) Gravação (transação única)
    with conn:
        if composta:
            conn.execute("DELETE FROM tabelas_compostas WHERE nome = ?", (indice,))
            _alertar(indice, f"tabela composta {indice} substituída pela tabela importada de {caminho.name}", caminho_alertas, resultado)
        if not info_existente or composta:
            registrar_indice(
                conn,
                InfoIndice(indice, nome or f"Tabela {indice}", fonte or f"Importado de {caminho.name}", tipo, None, url_fonte),
                commit=False,
            )
        elif nome or fonte or url_fonte:
            atual = conn.execute("SELECT nome, fonte, tipo, codigo_sgs, url_fonte FROM indices WHERE codigo = ?", (indice,)).fetchone()
            registrar_indice(
                conn,
                InfoIndice(indice, nome or atual[0], fonte or atual[1], atual[2], atual[3], url_fonte or atual[4]),
                commit=False,
            )
        for mes, taxa in fora:
            _alertar(indice, f"{mes:%m/%Y}: taxa de {_pct_br(taxa)} fora da faixa usual, aceita por opção explícita "
                     "(--aceitar-fora-da-faixa)", caminho_alertas, resultado)
        for mes, taxa, projecao in taxas:
            ja = existentes.get(mes)
            if ja is not None:
                if ja[0] == taxa and ja[1] == projecao:
                    continue
                if not ja[1] and projecao:
                    resultado.ignorados.append(f"{mes:%m/%Y}: projeção do arquivo não substitui valor oficial do banco")
                    continue
                if ja[1] and not projecao:
                    _alertar(indice, f"{mes:%m/%Y}: projeção {ja[0]} substituída pelo valor oficial {taxa}", caminho_alertas, resultado)
                elif not substituir and not (ja[1] and projecao):
                    continue
                resultado.substituidos += 1
            gravar_taxa(conn, indice, mes, taxa, projecao=projecao, substituir=True)
            resultado.gravados += 1
            resultado.primeiro_mes = min(resultado.primeiro_mes or mes, mes)
            resultado.ultimo_mes = max(resultado.ultimo_mes or mes, mes)
        # tabela importada pelo usuário: as atualizações do plugin nunca mexem nela
        conn.execute("INSERT OR REPLACE INTO procedencia_tabelas (codigo, origem) VALUES (?, 'usuario')", (indice,))
    recalcular_fatores(conn, indice)
    return resultado
