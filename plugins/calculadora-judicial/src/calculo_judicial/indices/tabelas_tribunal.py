"""Conversores das tabelas oficiais de tribunal para o CSV do importador.

Formatos tratados:

- TJSP — planilha "Atualização Monetária e Juros Moratórios" (.xlsm), aba ÍNDICES:
  coluna A = mês (dia 01), colunas com o fator acumulado de cada tabela prática
  (ex.: coluna O = "Tabela Prática TJSP - Lei 14.905/2024"). Convenção TJSP: o
  fator do mês M já embute o índice do mês M−1 ("dividir pelo fator do mês
  inicial e multiplicar pelo do mês final"), logo a taxa do mês M é
  ``fator(M+1)/fator(M) − 1`` → importar com ``convencao_fator='fator_inicio_do_mes'``.
- CJF — tabela gerada pelo SICOM (``tabelaCorMor.php``), arquivo "xls" que na
  verdade é HTML: uma linha por ano e 12 colunas de coeficientes multiplicadores
  ("valor em moeda da época × coeficiente do mês/ano"). O coeficiente do mês M
  inclui o índice do próprio mês M até o mês anterior à validade, logo a taxa do
  mês M é ``coef(M)/coef(M+1) − 1`` → importar com
  ``convencao_fator='coeficiente_multiplicador'``.

Os CSV gerados (``mes_referencia,fator``) ficam em ``dados/`` ao lado do arquivo
original, para rastreabilidade. Uso:

    python -m calculo_judicial.indices.tabelas_tribunal --tjsp-xlsm dados/x.xlsm --coluna O --saida dados/tjsp.csv
    python -m calculo_judicial.indices.tabelas_tribunal --cjf-html dados/x.html --saida dados/jf.csv
"""

from __future__ import annotations

import argparse
import csv
import re
from datetime import date, datetime
from decimal import Decimal
from html.parser import HTMLParser
from pathlib import Path

from calculo_judicial.decimais import para_decimal

MESES = 12


def _coluna_indice(letra: str) -> int:
    n = 0
    for ch in letra.strip().upper():
        n = n * 26 + (ord(ch) - ord("A") + 1)
    return n - 1


def ler_tjsp_xlsm(caminho: str | Path, coluna: str = "O", desde: date | None = date(1994, 7, 1)) -> list[tuple[date, Decimal]]:
    """Lê (mês, fator) da aba de índices da planilha do TJSP.

    ``coluna``: letra da coluna da tabela desejada (O = Lei 14.905/2024; N = antiga INPC;
    P = IPCA-E). Valores numéricos do Excel chegam como float e são convertidos por
    ``repr`` (representação mais curta, que reproduz o valor digitado com 6 casas).
    """
    from openpyxl import load_workbook

    wb = load_workbook(caminho, read_only=True, data_only=True, keep_vba=False)
    try:
        aba = next((ws for ws in wb.worksheets if "NDICES" in ws.title.upper()), wb.worksheets[-1])
        idx = _coluna_indice(coluna)
        saida: list[tuple[date, Decimal]] = []
        for linha in aba.iter_rows(values_only=True):
            if not linha or not isinstance(linha[0], datetime):
                continue
            mes = linha[0].date().replace(day=1)
            if desde and mes < desde:
                continue
            valor = linha[idx] if idx < len(linha) else None
            if valor is None or valor == "":
                continue
            texto = repr(valor) if isinstance(valor, float) else str(valor)
            saida.append((mes, para_decimal(texto)))
        return saida
    finally:
        wb.close()


class _ExtratorTabela(HTMLParser):
    def __init__(self) -> None:
        super().__init__()
        self.linhas: list[list[str]] = []
        self._linha: list[str] | None = None
        self._celula: list[str] | None = None

    def handle_starttag(self, tag, attrs):
        if tag == "tr":
            self._linha = []
        elif tag in ("td", "th") and self._linha is not None:
            self._celula = []

    def handle_endtag(self, tag):
        if tag in ("td", "th") and self._linha is not None and self._celula is not None:
            self._linha.append(" ".join("".join(self._celula).split()))
            self._celula = None
        elif tag == "tr" and self._linha is not None:
            self.linhas.append(self._linha)
            self._linha = None

    def handle_data(self, data):
        if self._celula is not None:
            self._celula.append(data)


def ler_cjf_html(caminho: str | Path, desde: date | None = date(1994, 7, 1)) -> list[tuple[date, Decimal]]:
    """Lê (mês, coeficiente) da tabela de correção do CJF (HTML/xls do SICOM).

    Considera só o primeiro bloco de linhas "ano + até 12 coeficientes" (o bloco
    seguinte, de SELIC acumulada em %, é ignorado).
    """
    bruto = Path(caminho).read_bytes()
    try:
        texto = bruto.decode("utf-8")
    except UnicodeDecodeError:
        texto = bruto.decode("iso-8859-1")
    extrator = _ExtratorTabela()
    extrator.feed(texto)
    saida: list[tuple[date, Decimal]] = []
    vistos: set[int] = set()
    for linha in extrator.linhas:
        celulas = [c for c in linha if c != ""]
        if not celulas or not re.fullmatch(r"(19|20)\d\d", celulas[0]):
            continue
        valores = celulas[1:]
        if not valores or any("%" in v for v in valores):
            continue  # bloco de SELIC acumulada
        ano = int(celulas[0])
        if ano in vistos:
            break  # segundo bloco com os mesmos anos
        vistos.add(ano)
        for mes_idx, valor in enumerate(valores[:MESES], start=1):
            mes = date(ano, mes_idx, 1)
            if desde and mes < desde:
                continue
            saida.append((mes, para_decimal(valor)))
    return saida


def gravar_csv(registros: list[tuple[date, Decimal]], destino: str | Path) -> Path:
    destino = Path(destino)
    destino.parent.mkdir(parents=True, exist_ok=True)
    with open(destino, "w", encoding="utf-8", newline="") as arq:
        escritor = csv.writer(arq)
        escritor.writerow(["mes_referencia", "fator"])
        for mes, fator in registros:
            escritor.writerow([mes.isoformat(), format(fator, "f")])
    return destino


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="Converte tabelas oficiais (TJSP xlsm, CJF html) em CSV")
    grupo = parser.add_mutually_exclusive_group(required=True)
    grupo.add_argument("--tjsp-xlsm")
    grupo.add_argument("--cjf-html")
    parser.add_argument("--coluna", default="O")
    parser.add_argument("--saida", required=True)
    args = parser.parse_args(argv)
    if args.tjsp_xlsm:
        registros = ler_tjsp_xlsm(args.tjsp_xlsm, args.coluna)
    else:
        registros = ler_cjf_html(args.cjf_html)
    destino = gravar_csv(registros, args.saida)
    print(f"{len(registros)} meses ({registros[0][0]:%m/%Y} a {registros[-1][0]:%m/%Y}) gravados em {destino}")
    return 0


if __name__ == "__main__":  # pragma: no cover
    raise SystemExit(main())
