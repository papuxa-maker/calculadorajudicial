"""Conversão entre o formulário da interface e o motor (funções puras, testáveis sem Streamlit).

Os campos de valor são texto (nunca float) no formato brasileiro: "1.234,56",
"10.000" (dez mil) ou "1234,56"; "1234.56" (ponto decimal, como exportam planilhas)
também é aceito quando não pode ser milhar. Datas só em dd/mm/aaaa (ou aaaa-mm-dd).
As tabelas exibidas na tela são as mesmas exportadas (``relatorio.tabelas``).
"""

from __future__ import annotations

import csv
import io
import json
import re
from datetime import date, datetime
from decimal import Decimal

from calculo_judicial.motor.parcela import NATUREZAS, TIPOS_JUROS, Abatimento, Parcela, Periodo
from calculo_judicial.relatorio.tabelas import TabelasMemoria, tabelas_da_memoria, total_exibido  # noqa: F401 - reexportação

COLUNAS_PARCELAS = ["descricao", "valor", "data_base", "natureza", "data_evento"]
COLUNAS_ABATIMENTOS = ["data", "valor", "descricao"]
# Todos os campos de Periodo que alteram o resultado: editar os períodos sem mudar nada reproduz o automático
COLUNAS_PERIODOS = [
    "de", "ate", "indice_correcao", "tipo_juros", "taxa_juros_am", "indice_juros", "capitalizacao_correcao",
    "capitalizacao_juros", "acrescimo_final_pct", "selic_sobre_juros",
]
COLUNAS_PERIODOS_OCULTAS = ["descricao", "base_legal", "gerado"]
BASE_LEGAL_EDITADO = "critério do título (período gerado pela regra e editado manualmente)"

_MILHAR = re.compile(r"^-?\d{1,3}(\.\d{3})+$")  # 10.000 / 1.234.567 (sem centavos)
_BR = re.compile(r"^-?(\d{1,3}(\.\d{3})+|\d+),\d+$")  # 1.234,56 / 1234,56
_INTEIRO = re.compile(r"^-?\d+$")
_PONTO_DECIMAL = re.compile(r"^-?\d+\.\d+$")  # 1234.56 (planilhas); "1.500" cai antes em _MILHAR
_DATA_BR = re.compile(r"^(\d{1,2})/(\d{1,2})/(\d{4})$")
_DATA_ISO = re.compile(r"^(\d{4})-(\d{2})-(\d{2})$")


class ErroFormulario(ValueError):
    def __init__(self, erros: list[str]):
        self.erros = erros
        super().__init__("; ".join(erros))


def _vazio(v) -> bool:
    return v is None or (isinstance(v, str) and not v.strip()) or (isinstance(v, float) and v != v)


def ler_valor(v, rotulo: str = "valor") -> Decimal | None:
    """Número digitado na interface (dinheiro ou percentual) em Decimal.

    "10.000" é dez mil (ponto de milhar brasileiro); "1.234,56" e "1234,56" têm
    centavos; "1234.56" só é lido com ponto decimal quando não pode ser milhar.
    Recusa NaN, infinito, notação científica e "1,234.56".
    """
    if _vazio(v):
        return None
    if isinstance(v, bool):
        raise ValueError(f"{rotulo}: não é número")
    if isinstance(v, Decimal):
        numero = v
    elif isinstance(v, int):
        numero = Decimal(v)
    elif isinstance(v, float):
        numero = Decimal(repr(v))  # tabelas do Streamlit/planilhas podem devolver float; repr preserva o digitado
    else:
        texto = str(v).strip().replace("R$", "").replace(" ", "").replace(" ", "")
        if _MILHAR.match(texto):
            numero = Decimal(texto.replace(".", ""))
        elif _BR.match(texto):
            numero = Decimal(texto.replace(".", "").replace(",", "."))
        elif _INTEIRO.match(texto) or _PONTO_DECIMAL.match(texto):
            numero = Decimal(texto)
        else:
            raise ValueError(f"{rotulo}: não é número válido: {v!r} (use 1.234,56)")
    if not numero.is_finite():
        raise ValueError(f"{rotulo}: não é número válido: {v!r}")
    return numero


def ler_percentual(v, rotulo: str, maximo: Decimal = Decimal(100)) -> Decimal | None:
    """Percentual digitado (ex.: "10" ou "12,5"), maior que zero e até ``maximo``."""
    numero = ler_valor(v, rotulo)
    if numero is None:
        return None
    if numero <= 0 or numero > maximo:
        raise ValueError(f"{rotulo}: informe um percentual maior que 0 e até {maximo}")
    return numero


def ler_data(v, rotulo: str = "data") -> date | None:
    """Data digitada: só dd/mm/aaaa (ou aaaa-mm-dd, como vem das planilhas). Recusa mês/ano e ano com 2 dígitos."""
    if _vazio(v):
        return None
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    if hasattr(v, "to_pydatetime"):  # pandas.Timestamp
        return v.to_pydatetime().date()
    texto = str(v).strip()
    m = _DATA_BR.match(texto)
    partes = (m.group(3), m.group(2), m.group(1)) if m else None
    if partes is None and len(texto) in (10, 19):
        m = _DATA_ISO.match(texto[:10])
        partes = m.groups() if m else None
    if partes is None:
        raise ValueError(f"{rotulo}: data inválida {texto!r} (use dd/mm/aaaa)")
    try:
        return date(int(partes[0]), int(partes[1]), int(partes[2]))
    except ValueError:
        raise ValueError(f"{rotulo}: data inexistente {texto!r}") from None


def parcelas_de_tabela(linhas: list[dict]) -> list[Parcela]:
    parcelas, erros = [], []
    for n, linha in enumerate(linhas, start=1):
        if all(_vazio(linha.get(c)) for c in ("descricao", "valor", "data_base")):
            continue
        try:
            valor = ler_valor(linha.get("valor"), "valor")
            data_base = ler_data(linha.get("data_base"), "data-base")
            if valor is None or data_base is None:
                raise ValueError("valor e data-base são obrigatórios")
            if valor <= 0:
                raise ValueError("o valor deve ser maior que zero")
            natureza = (linha.get("natureza") or "principal").strip()
            if natureza not in NATUREZAS:
                raise ValueError(f"natureza deve ser uma de {', '.join(NATUREZAS)}")
            parcelas.append(
                Parcela(
                    descricao=str(linha.get("descricao") or f"Parcela {n}").strip(),
                    valor_original=valor,
                    data_base=data_base,
                    natureza=natureza,
                    data_evento=ler_data(linha.get("data_evento"), "data do evento"),
                )
            )
        except (ValueError, TypeError) as erro:
            erros.append(f"parcela {n}: {erro}")
    if erros:
        raise ErroFormulario(erros)
    return parcelas


def abatimentos_de_tabela(linhas: list[dict]) -> list[Abatimento]:
    abatimentos, erros = [], []
    for n, linha in enumerate(linhas, start=1):
        if all(_vazio(linha.get(c)) for c in COLUNAS_ABATIMENTOS):
            continue
        try:
            valor = ler_valor(linha.get("valor"), "valor")
            data = ler_data(linha.get("data"), "data")
            if valor is None or data is None:
                raise ValueError("data e valor são obrigatórios")
            if valor <= 0:
                raise ValueError("o valor deve ser maior que zero")
            abatimentos.append(Abatimento(data, valor, str(linha.get("descricao") or f"Abatimento {n}").strip()))
        except (ValueError, TypeError) as erro:
            erros.append(f"abatimento {n}: {erro}")
    if erros:
        raise ErroFormulario(erros)
    return abatimentos


def _texto_campo(v) -> str:
    return "" if _vazio(v) else str(v).strip()


def _bool(v, padrao: bool = True) -> bool:
    if _vazio(v):
        return padrao
    if isinstance(v, str):
        return v.strip().lower() in ("true", "1", "sim", "s", "verdadeiro")
    return bool(v)


def assinatura_periodo(linha: dict) -> str:
    """Campos que alteram o resultado, normalizados (para saber se uma linha gerada pela regra foi editada)."""

    def norm(campo):
        v = linha.get(campo)
        try:
            if campo in ("de", "ate"):
                d = ler_data(v, campo)
                return d.isoformat() if d else ""
            if campo in ("taxa_juros_am", "acrescimo_final_pct"):
                n = ler_valor(v, campo)
                return format(n.normalize(), "f") if n is not None else ""
        except ValueError:
            return f"invalido:{v}"
        if campo == "selic_sobre_juros":
            return str(_bool(v))
        if campo == "capitalizacao_correcao":
            return _texto_campo(v) or "composta"
        if campo == "tipo_juros":
            return _texto_campo(v) or "nenhum"
        return _texto_campo(v)

    return json.dumps([norm(c) for c in COLUNAS_PERIODOS])


def periodos_de_tabela(linhas: list[dict]) -> list[Periodo]:
    periodos, erros = [], []
    for n, linha in enumerate(linhas, start=1):
        if all(_vazio(linha.get(c)) for c in ("de", "ate", "indice_correcao", "taxa_juros_am", "indice_juros")):
            continue
        try:
            de = ler_data(linha.get("de"), "de")
            ate = ler_data(linha.get("ate"), "até")
            if de is None or ate is None:
                raise ValueError("informe as datas 'de' e 'até'")
            tipo = _texto_campo(linha.get("tipo_juros")) or "nenhum"
            if tipo not in TIPOS_JUROS:
                raise ValueError(f"tipo de juros deve ser um de {', '.join(TIPOS_JUROS)}")
            descricao = _texto_campo(linha.get("descricao"))
            base_legal = _texto_campo(linha.get("base_legal"))
            gerado = _texto_campo(linha.get("gerado"))
            if not gerado:
                descricao, base_legal = "", ""  # linha digitada: o motor descreve o critério informado
            elif gerado != assinatura_periodo(linha):
                descricao, base_legal = "", BASE_LEGAL_EDITADO  # regra alterada à mão: descrever o que vale
            periodos.append(
                Periodo(
                    de=de,
                    ate=ate,
                    indice_correcao=_texto_campo(linha.get("indice_correcao")) or None,
                    tipo_juros=tipo,
                    taxa_juros_am=ler_valor(linha.get("taxa_juros_am"), "% a.m."),
                    indice_juros=_texto_campo(linha.get("indice_juros")) or None,
                    capitalizacao_correcao=_texto_campo(linha.get("capitalizacao_correcao"))
                    or ("simples" if _texto_campo(linha.get("indice_correcao")).upper() == "SELIC" else "composta"),
                    capitalizacao_juros=_texto_campo(linha.get("capitalizacao_juros")) or None,
                    acrescimo_final_pct=ler_valor(linha.get("acrescimo_final_pct"), "acréscimo final %"),
                    selic_sobre_juros=_bool(linha.get("selic_sobre_juros")),
                    descricao=descricao,
                    base_legal=base_legal,
                )
            )
        except (ValueError, TypeError) as erro:
            erros.append(f"período {n}: {erro}")
    if erros:
        raise ErroFormulario(erros)
    return periodos


def periodos_para_tabela(periodos: list[Periodo]) -> list[dict]:
    """Linhas da tabela editável com todos os campos do período e a assinatura do que foi gerado pela regra."""
    linhas = []
    for p in periodos:
        linha = {
            "de": p.de, "ate": p.ate, "indice_correcao": p.indice_correcao or "", "tipo_juros": p.tipo_juros,
            "taxa_juros_am": format(p.taxa_juros_am, "f") if p.taxa_juros_am is not None else "",
            "indice_juros": p.indice_juros or "", "capitalizacao_correcao": p.capitalizacao_correcao,
            "capitalizacao_juros": p.capitalizacao_juros or "",
            "acrescimo_final_pct": format(p.acrescimo_final_pct, "f") if p.acrescimo_final_pct is not None else "",
            "selic_sobre_juros": bool(p.selic_sobre_juros), "descricao": p.descricao, "base_legal": p.base_legal,
        }
        linha["gerado"] = assinatura_periodo(linha) if (p.descricao or p.base_legal) else ""
        linhas.append(linha)
    return linhas


def estender_ate(linhas: list[dict], data_antiga: date | None, data_nova: date) -> list[dict]:
    """Períodos que iam "até a data do cálculo" acompanham a nova data (duplicar ou mudar a data do cálculo)."""
    if data_antiga is None or data_antiga == data_nova:
        return linhas
    saida = []
    for linha in linhas:
        try:
            ate = ler_data(linha.get("ate"), "até")
        except ValueError:
            ate = None
        if ate != data_antiga:
            saida.append(linha)
            continue
        nova = dict(linha)
        nao_editada = bool(linha.get("gerado")) and linha.get("gerado") == assinatura_periodo(linha)
        nova["ate"] = data_nova if isinstance(linha.get("ate"), date) else data_nova.strftime("%d/%m/%Y")
        if nao_editada:
            nova["gerado"] = assinatura_periodo(nova)  # continua sendo o período da regra, não uma edição
        saida.append(nova)
    return saida


def ler_planilha_parcelas(nome: str, conteudo: bytes) -> list[dict]:
    """Lê CSV/XLSX com as colunas ``descricao,valor,data_base,natureza`` (SPEC, seção 7)."""
    if nome.lower().endswith((".xlsx", ".xlsm")):
        from openpyxl import load_workbook

        wb = load_workbook(io.BytesIO(conteudo), data_only=True)
        ws = wb.worksheets[0]
        linhas = list(ws.iter_rows(values_only=True))
        wb.close()
        if not linhas:
            return []
        cab = [str(c).strip().lower() if c is not None else "" for c in linhas[0]]
        saida = []
        for linha in linhas[1:]:
            registro = {}
            for k, v in zip(cab, linha):
                if isinstance(v, float):
                    v = repr(v).replace(".", ",")  # número da planilha: vírgula decimal, sem confusão com milhar
                elif hasattr(v, "isoformat"):
                    v = v.isoformat()[:10]
                registro[k] = "" if v is None else str(v)
            if any(registro.values()):
                saida.append(registro)
        return saida
    try:
        texto = conteudo.decode("utf-8-sig")
    except UnicodeDecodeError:
        texto = conteudo.decode("cp1252")
    primeira = texto.splitlines()[0] if texto else ""
    separador = ";" if primeira.count(";") > primeira.count(",") else ","
    leitor = csv.DictReader(io.StringIO(texto), delimiter=separador)
    return [{(k or "").strip().lower(): (v or "").strip() for k, v in linha.items()} for linha in leitor]
