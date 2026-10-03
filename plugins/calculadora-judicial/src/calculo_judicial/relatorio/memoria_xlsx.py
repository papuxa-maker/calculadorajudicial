"""Memória de cálculo em Excel (SPEC, seção 8) com openpyxl.

Abas: Resumo, Critérios, uma aba por parcela com a memória analítica (fórmulas
vivas: valor corrigido = valor base × fator acumulado; subtotal = corrigido +
juros; total da parcela = subtotal da última linha), Abatimentos e Fontes.
O total do Resumo é fórmula (soma dos totais das parcelas + multa + honorários),
para o Excel recalcular o mesmo total ao abrir. Formatação contábil R$ e datas
dd/mm/aaaa.
"""

from __future__ import annotations

import io
import math
import re
from datetime import date

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from calculo_judicial.datas import formatar_data
from calculo_judicial.decimais import formatar_brl
from calculo_judicial.indices.modelo import nome_curto
from calculo_judicial.motor.memoria import Memoria
from calculo_judicial.motor.parcela import NOMES_NATUREZA
from calculo_judicial.relatorio.tabelas import NOMES_IMPUTACAO

FORMATO_RS = '_-"R$" * #,##0.00_-;-"R$" * #,##0.00_-;_-"R$" * "-"??_-;_-@_-'
FORMATO_DATA = "dd/mm/yyyy"
FORMATO_COMPETENCIA = "mm/yyyy"
FORMATO_PCT = '0.0000"%"'
FORMATO_FATOR = "0.000000"
NEGRITO = Font(bold=True)
TITULO = Font(bold=True, size=14)
CABECALHO = PatternFill("solid", fgColor="E8EDF3")
DESTAQUE = PatternFill("solid", fgColor="FFF3CD")
FINO = Side(style="thin", color="B8C2CC")
BORDA = Border(left=FINO, right=FINO, top=FINO, bottom=FINO)


def _nome_aba(indice: int, descricao: str, usados: set[str]) -> str:
    base = re.sub(r"[\[\]\*\?/\\:']", " ", f"P{indice} {descricao}").strip()[:31]
    nome, n = base, 2
    while nome in usados:
        sufixo = f" ({n})"
        nome = base[: 31 - len(sufixo)] + sufixo
        n += 1
    usados.add(nome)
    return nome


def _cabecalho_linha(ws, linha: int, titulos: list[str]) -> None:
    for col, titulo in enumerate(titulos, start=1):
        c = ws.cell(row=linha, column=col, value=titulo)
        c.font = NEGRITO
        c.fill = CABECALHO
        c.border = BORDA
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _larguras(ws, larguras: list[float]) -> None:
    for i, w in enumerate(larguras, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def _aba_parcela(ws, mp) -> str:
    """Preenche a memória analítica da parcela; devolve a referência da célula com o total da parcela."""
    ws["A1"] = mp.parcela.descricao
    ws["A1"].font = TITULO
    ws["A2"] = (f"Valor original {formatar_brl(mp.valor_original)} — data-base {formatar_data(mp.parcela.data_base)} — "
                f"natureza {NOMES_NATUREZA.get(mp.parcela.natureza, mp.parcela.natureza)}")
    titulos = ["Competência", "Período", "Índice", "Valor base", "Índice %", "Fator acumulado", "Valor corrigido",
               "Juros %", "Juros acumulados", "Subtotal", "Observação"]
    _cabecalho_linha(ws, 4, titulos)
    linha = 4
    for l in mp.linhas:
        linha += 1
        ws.cell(row=linha, column=1, value=l.competencia).number_format = FORMATO_COMPETENCIA
        ws.cell(row=linha, column=2, value=l.periodo)
        ws.cell(row=linha, column=3, value=nome_curto(l.indice) + (f" ({nome_curto(l.indice_origem)})" if l.indice_origem and l.indice_origem != l.indice else ""))
        ws.cell(row=linha, column=4, value=l.base).number_format = FORMATO_RS
        c = ws.cell(row=linha, column=5, value=l.taxa_indice_pct)
        c.number_format = FORMATO_PCT
        ws.cell(row=linha, column=6, value=l.fator_acumulado).number_format = FORMATO_FATOR
        if l.tipo_juros == "ajuste":
            # linha de ajuste (truncamento CJF / preservação do nominal): valor, não fórmula
            ws.cell(row=linha, column=7, value=l.corrigido).number_format = FORMATO_RS
        else:
            ws.cell(row=linha, column=7, value=f"=D{linha}*F{linha}").number_format = FORMATO_RS
        ws.cell(row=linha, column=8, value=l.taxa_juros_pct).number_format = FORMATO_PCT
        ws.cell(row=linha, column=9, value=l.juros_acumulados).number_format = FORMATO_RS
        ws.cell(row=linha, column=10, value=f"=G{linha}+I{linha}").number_format = FORMATO_RS
        ws.cell(row=linha, column=11, value=("projeção; " if l.projecao else "") + l.observacao)
        if l.projecao or l.observacao:
            for col in range(1, 12):
                ws.cell(row=linha, column=col).fill = DESTAQUE
        for col in range(1, 11):
            ws.cell(row=linha, column=col).border = BORDA
    ultima = linha
    linha += 1
    ws.cell(row=linha, column=1, value="Total da parcela").font = NEGRITO
    if ultima >= 5:
        ws.cell(row=linha, column=7, value=f"=G{ultima}")
        ws.cell(row=linha, column=9, value=f"=I{ultima}")
        ws.cell(row=linha, column=10, value=f"=J{ultima}")
    else:
        ws.cell(row=linha, column=7, value=mp.principal_corrigido)
        ws.cell(row=linha, column=9, value=mp.juros)
        ws.cell(row=linha, column=10, value=mp.total)
    for col in (7, 9, 10):
        c = ws.cell(row=linha, column=col)
        c.number_format = FORMATO_RS
        c.font = NEGRITO
        c.border = BORDA
    ws.freeze_panes = "B5"
    _larguras(ws, [11, 8, 26, 15, 11, 14, 16, 11, 16, 16, 40])
    return f"'{ws.title}'!J{linha}"


def gerar_xlsx(memoria: Memoria, config=None, processo: dict | None = None) -> bytes:
    """Gera o Excel da memória de cálculo e devolve os bytes."""
    processo = processo or {}
    wb = Workbook()
    resumo = wb.active
    resumo.title = "Resumo"
    linha = 1
    if config and getattr(config, "nome_escritorio", ""):
        resumo.cell(row=linha, column=1, value=config.nome_escritorio).font = NEGRITO
        linha += 1
        if getattr(config, "oab", ""):
            resumo.cell(row=linha, column=1, value=config.oab)
            linha += 1
    resumo.cell(row=linha, column=1, value="Memória de Cálculo").font = TITULO
    linha += 1
    for rotulo, valor in (("Processo", processo.get("numero_cnj", "")), ("Vara/tribunal", processo.get("vara", "")),
                          ("Partes", f"{processo.get('autor', '')} × {processo.get('reu', '')}".strip(" ×")),
                          ("Tipo", processo.get("tipo", "")), ("Data do cálculo", memoria.parametros.data_calculo)):
        if valor:
            resumo.cell(row=linha, column=1, value=rotulo).font = NEGRITO
            c = resumo.cell(row=linha, column=2, value=valor)
            if rotulo == "Data do cálculo":
                c.number_format = FORMATO_DATA
                c.alignment = Alignment(horizontal="left")
            linha += 1
    linha += 1

    # Abas das parcelas primeiro (o total do Resumo referencia as células delas)
    usados = {"Resumo", "Critérios", "Abatimentos", "Fontes"}
    refs_totais = []
    for i, mp in enumerate(memoria.parcelas, start=1):
        ws = wb.create_sheet(_nome_aba(i, mp.parcela.descricao, usados))
        refs_totais.append(_aba_parcela(ws, mp))

    _cabecalho_linha(resumo, linha, ["Item", "Valor"])
    inicio = linha + 1
    r = memoria.resumo
    itens = [("Principal original", r.principal), ("Correção monetária", r.correcao), ("Juros de mora", r.juros),
             ("Multa", r.multa), ("Honorários", r.honorarios), ("Abatimentos", -r.abatimentos)]
    for nome, valor in itens:
        linha += 1
        resumo.cell(row=linha, column=1, value=nome).border = BORDA
        c = resumo.cell(row=linha, column=2, value=valor)
        c.number_format = FORMATO_RS
        c.border = BORDA
    linha_multa, linha_honorarios = inicio + 3, inicio + 4
    linha += 1
    resumo.cell(row=linha, column=1, value="Total devido").font = NEGRITO
    celula_total = resumo.cell(row=linha, column=2)  # fórmula escrita depois do quadro "Total de cada parcela"
    celula_total.number_format = FORMATO_RS
    celula_total.font = NEGRITO
    celula_total.border = BORDA
    linha_total = linha
    linha += 1
    resumo.cell(row=linha, column=1, value="Conferência: soma das linhas acima")
    resumo.cell(row=linha, column=2, value=f"=SUM(B{inicio}:B{linha_total - 1})").number_format = FORMATO_RS
    linha += 2
    resumo.cell(row=linha, column=1, value="Detalhe dos acessórios").font = NEGRITO
    ac = memoria.acessorios
    pr = memoria.parametros
    for nome, valor, pedido in (("Multa contratual", ac.multa_contratual, pr.multa_contratual_pct is not None),
                                ("Honorários de sucumbência", ac.honorarios_sucumbencia, pr.honorarios_pct is not None),
                                ("Multa do art. 523", ac.multa_523, pr.multa_523),
                                ("Honorários do cumprimento (art. 523)", ac.honorarios_cumprimento, pr.honorarios_cumprimento)):
        if not (pedido or valor):
            continue
        linha += 1
        resumo.cell(row=linha, column=1, value=nome)
        resumo.cell(row=linha, column=2, value=valor).number_format = FORMATO_RS
    if r.saldo_credor:
        linha += 1
        resumo.cell(row=linha, column=1, value="Saldo credor (pagamento acima do devido)")
        resumo.cell(row=linha, column=2, value=r.saldo_credor).number_format = FORMATO_RS
    from calculo_judicial.relatorio.memoria_pdf import notas_do_resumo

    notas = notas_do_resumo(memoria)
    if notas:
        linha += 1
    for nota in notas:
        linha += 1
        c = resumo.cell(row=linha, column=1, value=nota)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        resumo.merge_cells(start_row=linha, start_column=1, end_row=linha, end_column=2)
        resumo.row_dimensions[linha].height = 15 * math.ceil(len(nota) / 60)  # célula mesclada não se ajusta sozinha

    # Total de cada parcela, lido das abas: o Total devido soma este quadro (uma fórmula por parcela, e não uma
    # fórmula gigante — o Excel recusa fórmulas acima de 8.192 caracteres). Honorários fixados em quantia certa já
    # estão na linha "Honorários" e ficam fora da soma.
    linha += 2
    resumo.cell(row=linha, column=1, value="Total de cada parcela (das abas)").font = NEGRITO
    credito = [(ref, mp) for ref, mp in zip(refs_totais, memoria.parcelas) if mp.parcela.natureza != "honorarios"]
    honorarios = [(ref, mp) for ref, mp in zip(refs_totais, memoria.parcelas) if mp.parcela.natureza == "honorarios"]
    primeira = linha + 1
    for ref, mp in credito:
        linha += 1
        resumo.cell(row=linha, column=1, value=mp.parcela.descricao)
        resumo.cell(row=linha, column=2, value=f"={ref}").number_format = FORMATO_RS
    ultima = linha
    for ref, mp in honorarios:
        linha += 1
        resumo.cell(row=linha, column=1, value=f"{mp.parcela.descricao} (já na linha Honorários)")
        resumo.cell(row=linha, column=2, value=f"={ref}").number_format = FORMATO_RS
    termos = ([f"SUM(B{primeira}:B{ultima})"] if credito else []) + [f"B{linha_multa}", f"B{linha_honorarios}"]
    celula_total.value = "=" + "+".join(termos)
    _larguras(resumo, [42, 22])

    # Critérios
    crit = wb.create_sheet("Critérios", 1)
    crit["A1"] = "Critérios adotados"
    crit["A1"].font = TITULO
    linha = 2
    for texto in memoria.criterios + memoria.acessorios.detalhes:
        linha += 1
        crit.cell(row=linha, column=1, value=texto).alignment = Alignment(wrap_text=True, vertical="top")
    if memoria.avisos:
        linha += 2
        crit.cell(row=linha, column=1, value="Observações").font = NEGRITO
        for aviso in memoria.avisos:
            linha += 1
            crit.cell(row=linha, column=1, value=f"• {aviso}").alignment = Alignment(wrap_text=True, vertical="top")
    linha += 2
    _cabecalho_linha(crit, linha, ["Parâmetro", "Valor"])
    for p in (p for p in memoria.parametros_relevantes if p.aplica):
        linha += 1
        crit.cell(row=linha, column=1, value=p.descricao)
        crit.cell(row=linha, column=2, value=p.valor)
    _larguras(crit, [110, 40])

    # Abatimentos
    ab = wb.create_sheet("Abatimentos")
    _cabecalho_linha(ab, 1, ["Data", "Descrição", "Saldo antes", "Valor abatido", "Juros imputados", "Capital imputado",
                             "Imputação", "Saldo após", "Excedente"])
    for i, a in enumerate(memoria.abatimentos, start=2):
        ab.cell(row=i, column=1, value=a.data).number_format = FORMATO_DATA
        ab.cell(row=i, column=2, value=a.descricao)
        for col, valor in ((3, a.saldo_antes), (4, a.valor), (5, a.juros_imputados), (6, a.capital_imputado), (8, a.saldo_apos),
                           (9, a.excedente)):
            ab.cell(row=i, column=col, value=valor).number_format = FORMATO_RS
        ab.cell(row=i, column=7, value=NOMES_IMPUTACAO.get(a.imputacao, a.imputacao.replace("_", " ")))
    if not memoria.abatimentos:
        ab.cell(row=2, column=1, value="Nenhum pagamento ou abatimento informado.")
    _larguras(ab, [12, 30, 16, 16, 16, 16, 18, 16, 14])

    # Fontes
    fo = wb.create_sheet("Fontes")
    _cabecalho_linha(fo, 1, ["Índice", "Nome", "Órgão/fonte", "URL", "Data de coleta", "Primeiro mês", "Último mês", "Projeção"])
    for i, f in enumerate(memoria.fontes, start=2):
        fo.cell(row=i, column=1, value=nome_curto(f.codigo))
        fo.cell(row=i, column=2, value=f.nome)
        fo.cell(row=i, column=3, value=f.fonte)
        fo.cell(row=i, column=4, value=f.url or "")
        try:
            fo.cell(row=i, column=5, value=date.fromisoformat(f.data_coleta[:10])).number_format = FORMATO_DATA
        except ValueError:
            fo.cell(row=i, column=5, value=f.data_coleta)
        if f.primeiro_mes:
            fo.cell(row=i, column=6, value=f.primeiro_mes).number_format = FORMATO_COMPETENCIA
            fo.cell(row=i, column=7, value=f.ultimo_mes).number_format = FORMATO_COMPETENCIA
        fo.cell(row=i, column=8, value=", ".join(m.strftime("%m/%Y") for m in f.meses_projecao))
        if f.meses_projecao:
            for col in range(1, 9):
                fo.cell(row=i, column=col).fill = DESTAQUE
    _larguras(fo, [14, 40, 30, 50, 14, 12, 12, 20])

    # ordem das abas: Resumo, Critérios, parcelas..., Abatimentos, Fontes
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
